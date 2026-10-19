"""Build the first Word file of the corpus from the learner's own interview notes (step 2.7).

Why a generated file: open-licence .docx files are rare online, and the learner chose
their own notes (own work = public, corpus group C, D-043). The Markdown in docs/interview/
stays the source of truth; this script turns it into a real .docx, like notes written in
Word, plus one table of real C0 results (copied from evals/ABLATION.md) so Docling has a
table to read.

Run: python scripts/make_notes_docx.py   -> data/docs/interview/phase_1_notes.docx
"""

import re
from pathlib import Path

from docx import Document  # python-docx, installed with docling-slim[format-docx]

from docpilot import config

SOURCE = config.PROJECT_ROOT / "docs" / "interview" / "phase_1_interview.md"
OUTPUT = config.DOCS_DIR / "interview" / "phase_1_notes.docx"
HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
BULLET = re.compile(r"^\s*[-*]\s+(.*)$")

# Real numbers from evals/ABLATION.md (C0, dev split, golden v0, n=19; latency n=22).
RESULTS = [
    ("Metric", "What it measures", "C0 value"),
    ("Recall@1", "right chunk ranked first", "0.526"),
    ("Recall@5", "right chunk somewhere in the top 5", "0.842"),
    ("MRR", "how high the first right chunk is ranked", "0.667"),
    ("nDCG@5", "all right chunks near the top", "0.670"),
    ("Abstention accuracy", "refuses exactly when the answer is not in the docs", "0.909"),
    ("p50 latency", "typical end-to-end answer time, warm", "1107 ms"),
    ("p95 latency", "slow-tail answer time, warm", "1764 ms"),
]


def add_text(paragraph, text: str) -> None:
    """Add text with **bold** parts as bold runs; backticks are dropped (Word has none)."""
    for index, part in enumerate(text.replace("`", "").split("**")):
        if part:
            paragraph.add_run(part).bold = index % 2 == 1  # odd parts were between ** **


def build(source: Path, output: Path) -> None:
    document = Document()
    paragraph_lines: list[str] = []

    def flush() -> None:
        if paragraph_lines:
            add_text(document.add_paragraph(), " ".join(paragraph_lines))
            paragraph_lines.clear()

    for line in source.read_text(encoding="utf-8").splitlines():
        heading, bullet = HEADING.match(line), BULLET.match(line)
        if not line.strip() or line.strip() == "---":
            flush()  # blank line or horizontal rule ends a paragraph
        elif heading:
            flush()
            document.add_heading(heading.group(2).replace("`", ""), level=len(heading.group(1)))
        elif bullet:
            flush()
            add_text(document.add_paragraph(style="List Bullet"), bullet.group(1))
        else:
            paragraph_lines.append(line.strip())  # wrapped lines of one paragraph
    flush()

    document.add_heading("Phase 1 results (C0 baseline, dev split)", level=2)
    document.add_paragraph(
        "Naive baseline C0: fixed-size chunks of 800 characters with 100 overlap, dense "
        "search only, top 4 chunks to the LLM. 19 answerable dev questions."
    )
    table = document.add_table(rows=len(RESULTS), cols=3)
    table.style = "Table Grid"
    for row, values in zip(table.rows, RESULTS, strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"saved {output} ({output.stat().st_size} bytes)")


if __name__ == "__main__":
    build(SOURCE, OUTPUT)
