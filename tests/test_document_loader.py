"""Tests for ingestion/document_loader.py (Docling) on a small Word file made in the test.

DOCX needs no layout models, so this runs offline in CI. The PDF path was checked by hand
on the RAGAS paper (8 pages, 4 tables, 26 s warm; docs/2_decisions.md D-043).
"""

import pytest
from docx import Document

from docpilot.ingestion import document_loader
from docpilot.ingestion.document_loader import convert_to_markdown
from docpilot.ingestion.loader import load_documents
from docpilot.ingestion.sources import SOURCES


def make_docx(path, value: str = "0.667") -> None:
    document = Document()
    document.add_heading("Phase 1 results", level=1)
    document.add_paragraph("Numbers from the baseline.")
    table = document.add_table(rows=2, cols=2)
    for cell, text in zip(
        table.rows[0].cells + table.rows[1].cells, ["Metric", "Value", "MRR", value], strict=True
    ):
        cell.text = text
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def test_word_table_becomes_a_markdown_table(tmp_path):
    path = tmp_path / "docs" / "notes" / "results.docx"
    make_docx(path)
    markdown = convert_to_markdown(path, "notes/results.docx", cache_dir=tmp_path / "cache")
    assert "Phase 1 results" in markdown
    table_rows = [line for line in markdown.splitlines() if line.startswith("|")]
    assert any("MRR" in row and "0.667" in row for row in table_rows)


def test_second_conversion_comes_from_the_cache(tmp_path, monkeypatch):
    path = tmp_path / "results.docx"
    make_docx(path)
    first = convert_to_markdown(path, "results.docx", cache_dir=tmp_path / "cache")

    def no_docling():
        raise AssertionError("Docling was called although nothing changed")

    monkeypatch.setattr(document_loader, "get_converter", no_docling)
    assert convert_to_markdown(path, "results.docx", cache_dir=tmp_path / "cache") == first


def test_changed_file_is_converted_again(tmp_path):
    path = tmp_path / "results.docx"
    make_docx(path, value="0.667")
    convert_to_markdown(path, "results.docx", cache_dir=tmp_path / "cache")
    make_docx(path, value="0.795")  # same name, new content -> new sha256
    assert "0.795" in convert_to_markdown(path, "results.docx", cache_dir=tmp_path / "cache")


def test_loader_reads_docx_next_to_markdown(tmp_path, monkeypatch):
    monkeypatch.setattr(document_loader.config, "CONVERTED_DIR", tmp_path / "cache")
    (tmp_path / "docs" / "fastapi").mkdir(parents=True)
    (tmp_path / "docs" / "fastapi" / "a.md").write_text("# A\n", encoding="utf-8")
    make_docx(tmp_path / "docs" / "notes" / "b.docx")
    documents = load_documents(tmp_path / "docs")
    assert [doc.source for doc in documents] == ["fastapi/a.md", "notes/b.docx"]
    assert "MRR" in documents[1].text


@pytest.mark.parametrize("name", ["ragas", "interview"])
def test_new_sources_have_licence_and_visibility(name):
    info = SOURCES[name]
    assert info.licence and info.visibility in ("public", "local_only")
